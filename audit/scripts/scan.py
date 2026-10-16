#!/usr/bin/env python3
"""Vendor-mention dictionary scan over every tracked file of learn-ccna1.

Usage: scan.py [PATH ...]
  no PATH   - full scan (input for build.py)
  PATH ...  - only tracked files under these paths (check a work package)
Output in $AUDIT_WORK/text: hits.jsonl (one record per dictionary match) + files.jsonl (coverage per file).
The audit/ directory itself is never scanned (it quotes every term it reports)."""
import re, os, sys, json, subprocess, zipfile, collections, tempfile
HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.environ.get('AUDIT_REPO') or os.path.abspath(os.path.join(HERE, '..', '..'))
WORK = os.environ.get('AUDIT_WORK') or os.path.join(tempfile.gettempdir(), 'vendor-audit')
OUT = os.path.join(WORK, 'text'); os.makedirs(OUT, exist_ok=True)
PREFIXES = [p.rstrip('/') for p in sys.argv[1:]]
TEXT_EXT = {'.md', '.json', '.yaml', '.yml', '.tf', '.sh', '.tpl', '.py', '.txt', '.gitignore'}
IMG_EXT = {'.png', '.jpg', '.jpeg', '.gif', '.svg'}
BIN_EXT = {'.pka', '.pkt'}
I = re.IGNORECASE
# key, regex, flags, default category, in-issue-dictionary?
TERMS = [
 ('cisco', r'cisco', I, 'brand', True),
 ('сиско/циско', r'сиско|циско', I, 'brand', True),
 ('netacad', r'netacad', I, 'brand', True),
 ('networking academy', r'networking\s+academy', I, 'brand', True),
 ('сетевая академия', r'сетев\w*\s+академи\w*', I, 'brand', True),
 ('академи', r'академи\w*', I, 'brand', True),
 ('skills for all', r'skills\s*for\s*all|skillsforall', I, 'brand', False),
 ('learning network', r'learning\s*network', I, 'brand', False),
 ('talos', r'\btalos\b', I, 'brand', False),
 ('ccna', r'ccna', I, 'certification', True),
 ('ccnp', r'ccnp', I, 'certification', True),
 ('ccent', r'ccent', I, 'certification', True),
 ('ccie', r'ccie', I, 'certification', False),
 ('cyberops', r'cyberops', I, 'certification', False),
 ('itn', r'\bitn\b', I, 'certification', True),
 ('srwe', r'\bsrwe\b', I, 'certification', True),
 ('ensa', r'\bensa\b', I, 'certification', True),
 ('v7', r'\bv7(?:\.0\d)?\b', I, 'certification', True),
 ('scope and sequence', r'scope\s*(?:and|&)\s*sequence', I, 'certification', False),
 ('сертификаци', r'сертификаци\w*|сертифицир\w*', I, 'certification', False),
 ('packet tracer', r'packet\s*tracer|packettracer', I, 'product', True),
 ('PT', r'\bPT\b', 0, 'product', True),
 ('PTSA/PTPM', r'\bPT(?:SA|PM)\b', 0, 'lab-infra', False),
 ('.pka', r'\.pka\b', I, 'lab-infra', True),
 ('.pkt', r'\.pkt\b', I, 'lab-infra', True),
 ('webex', r'webex', I, 'product', True),
 ('meraki', r'meraki', I, 'product', True),
 ('catalyst', r'catalyst', I, 'product', True),
 ('ISR', r'\bISR\b', 0, 'product', True),
 ('nexus', r'\bnexus\b', I, 'product', True),
 ('ASA', r'\bASA\b', 0, 'product', True),
 ('anyconnect', r'anyconnect', I, 'product', True),
 ('dna center', r'dna\s*cent(?:er|re)', I, 'product', True),
 ('IOS', r'\bIOS\b(?![\s-]*XE)', I, 'product', True),
 ('IOS XE', r'\bIOS[\s-]*XE\b', I, 'product', True),
 ('NX-OS', r'\bNX-?OS\b', I, 'product', True),
 ('umbrella', r'\bumbrella\b', I, 'product', False),
 ('linksys', r'linksys', I, 'product', False),
 ('aironet', r'aironet', I, 'product', False),
 ('firepower', r'firepower', I, 'product', False),
 ('ios-image', r'\b[ci]\d{4}\S*k9\S*|universalk9|lanbasek9|lanbase|ipbasek9|ipservicesk9|\bmz\.\d{3}-', I, 'model', False),
 ('model-number', r'\b(?:1941|2901|2911|2951|2960|3560|3650|3750|4221|4321|4331|4451|1841|2811|1921|9200|9300|3850|1801|1811|1812|1861|2801|2821|2851|3825|3845|881|887|891)\b', 0, 'model', True),
 ('WS-C', r'WS-C\S*', 0, 'model', True),
 ('ios-version', r'\b1[25]\.\d\(\d+[a-z]?\)[A-Za-z0-9]*|\b12\.2\.\d+-SE\d*|\b15\.0\.2-SE\d*|(?<=XE )(?:Release |Version |версии )?16\.\d{1,2}\.\d{1,2}|(?<=XE версии )16\.\d{1,2}\.\d{1,2}|(?<=Version )16\.\d{2}\.\d{2}', 0, 'model', False),
 ('PLE400', r'PLEK?400', 0, 'model', False),
 ('MX64W', r'MX64W?', 0, 'model', False),
 ('series', r'\b(?:1900|2900|4000|800|1800|3800|4300|2800|4400|3900)\s*(?:series|серии)\b', I, 'model', False),
 ('copyright', r'copyright|all rights reserved', I, 'legal', True),
 ('©', r'©|\(c\)\s*(?:19|20)\d\d', I, 'legal', True),
 ('®', r'®', 0, 'legal', True),
 ('™', r'™', 0, 'legal', True),
 ('товарн', r'товарн\w*', I, 'legal', True),
 ('trademark', r'trademark', I, 'legal', False),
 ('CDP', r'\bCDP\b', 0, 'proprietary-tech', True),
 ('EIGRP', r'EIGRP', I, 'proprietary-tech', True),
 ('HSRP', r'HSRP', I, 'proprietary-tech', True),
 ('GLBP', r'GLBP', I, 'proprietary-tech', True),
 ('VTP', r'\bVTP\b', 0, 'proprietary-tech', True),
 ('DTP', r'\bDTP\b', 0, 'proprietary-tech', True),
 ('ISL', r'\bISL\b', 0, 'proprietary-tech', True),
 ('PVST', r'PVST', I, 'proprietary-tech', True),
 ('SD-Access', r'SD-Access', I, 'proprietary-tech', False),
 ('LWAPP', r'LWAPP', I, 'proprietary-tech', False),
 ('PAgP', r'\bPAgP\b', I, 'proprietary-tech', False),
 ('EtherChannel', r'etherchannel', I, 'proprietary-tech', False),
 ('NetFlow', r'netflow', I, 'proprietary-tech', False),
 ('CEF', r'\bCEF\b|cisco express forwarding', 0, 'proprietary-tech', False),
 ('PortFast', r'portfast', I, 'proprietary-tech', False),
 ('BPDU Guard', r'bpdu\s*guard', I, 'proprietary-tech', False),
 ('IGRP', r'\bIGRP\b', 0, 'proprietary-tech', False),
 ('IP SLA', r'\bIP\s+SLA\b', 0, 'proprietary-tech', False),
 ('DMVPN', r'DMVPN', I, 'proprietary-tech', False),
 ('AutoSecure', r'autosecure|auto\s+secure', I, 'proprietary-tech', False),
 ('cisco.com', r'[\w.-]*cisco\.com', I, 'link', True),
 ('netacad.com', r'netacad\.com', I, 'link', False),
 ('youtube', r'youtube\.com/watch\?v=[\w-]+|youtu\.be/[\w-]+', I, 'link', False),
 ('pt-lab', r'pt-lab|pt_lab|pt-url|pt-api|validate_web_response|packettracer\.exe', I, 'lab-infra', False),
 ('db8:acad', r'db8:acad', I, 'brand', False),
 ('ACAD_', r'\bACAD_[A-Z]+', 0, 'brand', False),
 ('opendns', r'opendns', I, 'product', False),
 ('telepresence', r'telepresence', I, 'product', False),
 ('ESA/WSA', r'\bESA\b|\bWSA\b', 0, 'product', False),
 ('ACI/APIC', r'\bACI\b|\bAPIC\b', 0, 'product', False),
 ('ASDM', r'\bASDM\b', 0, 'product', False),
]
CTERMS = [(k, re.compile(rx, fl), cat, indict) for k, rx, fl, cat, indict in TERMS]

def module_of(p):
    top = p.split('/')[0]
    if re.match(r'^\d\d-', top) or top.startswith('final-') or top == '__exams': return top
    return '(root)'
def material_of(p):
    parts = p.split('/')
    top = parts[0]
    if re.match(r'^\d\d-', top) and len(parts) > 2: return '/'.join(parts[:2])
    if top == '__exams' and len(parts) > 2: return '/'.join(parts[:3]) if len(parts) > 3 else '/'.join(parts[:2])
    if top.startswith('final-'): return top
    if re.match(r'^\d\d-', top): return top
    return parts[0] if len(parts) == 1 else '/'.join(parts[:1])

def surface_of(p):
    parts = p.split('/'); top = parts[0]; base = parts[-1]
    if top in ('videos',): return 'internal-doc', 'internal'
    if top in ('.github', '.vscode', 'animations', 'readme-assets', 'packet-tracer-source-files') or base in ('figma-assets.json', 'migration-instructions.md', 'ccna1-labs.md', 'labsComments.md', 'calc_duration.txt', '.gitignore'):
        return 'internal-doc', 'internal'
    if '/deploy/' in p or '/deploy-azure/' in p or '/assessment/' in p or base == 'assessment.json' or base == 'test.json':
        return 'lab-infra', 'internal'
    if top == '__exams': return 'quiz', 'internal'
    if base == 'learn-metadata.json' or base == 'material-tags.json': return 'metadata', 'learner'
    if base == 'README.md': return 'text', 'learner'
    if p.endswith('.pdf'): return 'text', 'learner'
    return 'text', 'learner'

def is_quiz_meta(p, content):
    return '"questions"' in content

def scan_line(line):
    """return list of (key, cat, indict, start, end, matched)"""
    out = []
    for k, rx, cat, indict in CTERMS:
        for m in rx.finditer(line):
            out.append((k, cat, indict, m.start(), m.end(), m.group(0)))
    return out

def iter_md_lines(text):
    """yield (lineno, line, state) where state in text/code/comment"""
    in_code = False; in_comment = False; fence = None
    for i, line in enumerate(text.split('\n'), 1):
        state_comment = in_comment
        s = line.strip()
        # comments can start mid-line; treat a line as comment if inside or if it opens one
        opens = '<!--' in line; closes = '-->' in line
        if in_comment:
            st = 'comment'
            if closes and line.rfind('-->') >= line.rfind('<!--'): in_comment = False
        elif opens:
            st = 'comment'
            if not (closes and line.rfind('-->') > line.rfind('<!--')): in_comment = True
        else:
            if re.match(r'^\s*(```|~~~)', line):
                if not in_code: in_code = True; fence = s[:3]
                elif s.startswith(fence): in_code = False
                st = 'code'
            else:
                st = 'code' if in_code else 'text'
        yield i, line, st

files = subprocess.run(['git', 'ls-files'], capture_output=True, text=True, cwd=REPO).stdout.split('\n')
files = [f for f in files if f and not f.startswith('audit/')]
if PREFIXES:
    files = [f for f in files if any(f == p or f.startswith(p + '/') for p in PREFIXES)]
hits = []; cover = []
def add(p, loc, line_text, st, surface, audience, extra=None):
    ms = scan_line(line_text)
    for k, cat, indict, a, b, mt in ms:
        rec = dict(path=p, locator=loc, key=k, cat=cat, indict=indict, start=a, end=b, matched=mt,
                   context=line_text.strip()[:300], state=st, surface=surface, audience=audience,
                   module=module_of(p), material=material_of(p))
        if extra: rec.update(extra)
        hits.append(rec)
    return len(ms)

for p in files:
    ext = os.path.splitext(p)[1].lower()
    full = os.path.join(REPO, p)
    surface, audience = surface_of(p)
    n = 0; how = ''
    if ext in TEXT_EXT or p.endswith('.gitignore'):
        text = open(full, encoding='utf-8', errors='replace').read()
        if p.endswith('.md'):
            for i, line, st in iter_md_lines(text):
                sf = surface
                if surface == 'text': sf = {'text': 'text', 'code': 'code', 'comment': 'comment'}[st]
                aud = audience if st != 'comment' else 'internal'
                n += add(p, f'{p}:{i}', line, st, sf, aud)
            how = 'md-lines'
        else:
            sf = surface
            if p.endswith('learn-metadata.json') and is_quiz_meta(p, text): sf = 'quiz'
            if p.endswith('test.json'): sf = 'quiz'
            for i, line in enumerate(text.split('\n'), 1):
                n += add(p, f'{p}:{i}', line, 'text', sf, audience)
            how = 'lines'
    elif ext == '.docx':
        try:
            import docx
            d = docx.Document(full)
            paras = [x.text for x in d.paragraphs]
            for t in d.tables:
                for r in t.rows:
                    for c in r.cells: paras.append(c.text)
            for i, t in enumerate(paras, 1):
                if t.strip(): n += add(p, f'{p}:para{i}', t, 'text', 'internal-doc' if 'packet-tracer-source-files' in p or '__exams' in p else 'text', 'internal' if ('packet-tracer-source-files' in p or '__exams' in p) else 'learner')
            how = f'docx-paragraphs({len(paras)})'
        except Exception as e:
            how = f'docx-ERROR {e}'
    elif ext == '.pdf':
        r = subprocess.run(['pdftotext', '-layout', full, '-'], capture_output=True, text=True)
        pages = r.stdout.split('\f')
        aud = 'internal' if p.startswith('readme-assets') else 'learner'
        sf = 'internal-doc' if p.startswith('readme-assets') else 'text'
        for pi, pg in enumerate(pages, 1):
            for li, line in enumerate(pg.split('\n'), 1):
                n += add(p, f'{p}:p{pi}:l{li}', line, 'text', sf, aud)
        how = f'pdftotext({len(pages)} pages)'
    elif ext in IMG_EXT:
        how = 'image → visual review (see render/raster sheets)'
    elif ext in BIN_EXT:
        how = 'binary PT file → file-level lab-infra row'
    else:
        how = 'other'
    cover.append(dict(path=p, ext=ext, how=how, hits=n, surface=surface))

# file & directory names
seen_names = set()
for p in files:
    parts = p.split('/')
    for i in range(len(parts)):
        comp = '/'.join(parts[:i+1])
        if comp in seen_names: continue
        seen_names.add(comp)
        name = parts[i]
        ms = scan_line(name)
        for k, cat, indict, a, b, mt in ms:
            hits.append(dict(path=comp, locator=comp, key=k, cat=cat, indict=indict, start=a, end=b, matched=mt,
                             context=name, state='name', surface='filename',
                             audience='internal' if (parts[0] in ('__exams','packet-tracer-source-files','readme-assets','videos','.github','animations') or '/deploy/' in comp or '/assessment/' in comp) else 'learner',
                             module=module_of(comp), material=material_of(comp)))

with open(os.path.join(OUT, 'hits.jsonl'), 'w') as f:
    for h in hits: f.write(json.dumps(h, ensure_ascii=False) + '\n')
with open(os.path.join(OUT, 'files.jsonl'), 'w') as f:
    for c in cover: f.write(json.dumps(c, ensure_ascii=False) + '\n')
print('files', len(files), 'hits', len(hits))
c = collections.Counter(h['key'] for h in hits)
for k, v in c.most_common(): print(f'{v:6d}  {k}')
print('--- by surface'); 
for k, v in collections.Counter(h['surface'] for h in hits).most_common(): print(f'{v:6d}  {k}')
print('--- coverage how');
for k, v in collections.Counter(c['how'].split('(')[0] for c in cover).most_common(): print(f'{v:6d}  {k}')
